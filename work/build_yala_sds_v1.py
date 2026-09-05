from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\IT4-K63\UDPTDN\youtube-ai-learning-assistant")
OUT = ROOT / "docs" / "architecture" / "SDS_YouTube_AI_Learning_Assistant_V1.docx"
ASSETS = ROOT / "work" / "yala_sds_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
NAVY = "1F4E78"
PALE = "EAF1F8"
ALT = "F5F7FA"
BORDER = "D9D9D9"
GRAY = "666666"
FONT = "Arial"


def set_font(run, size=None, bold=None, color=BLACK, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), BORDER)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, 9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Cập nhật mục lục trong Microsoft Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def set_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def font_path(bold=False):
    return r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"


def pil_font(size, bold=False):
    return ImageFont.truetype(font_path(bold), size)


def text_center(draw, rect, title, subtitle=""):
    x1, y1, x2, y2 = rect
    title_font = pil_font(28, True)
    sub_font = pil_font(20)
    title_box = draw.multiline_textbbox((0, 0), title, font=title_font, align="center", spacing=5)
    th = title_box[3] - title_box[1]
    sh = 0
    if subtitle:
        sub_box = draw.multiline_textbbox((0, 0), subtitle, font=sub_font, align="center", spacing=4)
        sh = sub_box[3] - sub_box[1]
    total = th + (16 if subtitle else 0) + sh
    ty = y1 + ((y2 - y1) - total) / 2
    draw.multiline_text(((x1 + x2) / 2, ty), title, font=title_font, fill="#111111", anchor="ma", align="center", spacing=5)
    if subtitle:
        draw.multiline_text(((x1 + x2) / 2, ty + th + 16), subtitle, font=sub_font, fill="#444444", anchor="ma", align="center", spacing=4)


def box(draw, rect, title, subtitle="", fill="#F3F6FA"):
    draw.rounded_rectangle(rect, radius=18, fill=fill, outline="#1F4E78", width=3)
    text_center(draw, rect, title, subtitle)


def arrow(draw, start, end, label=""):
    draw.line([start, end], fill="#555555", width=4)
    x2, y2 = end
    x1, y1 = start
    import math
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 15
    left = (x2 - size * math.cos(angle - 0.6), y2 - size * math.sin(angle - 0.6))
    right = (x2 - size * math.cos(angle + 0.6), y2 - size * math.sin(angle + 0.6))
    draw.polygon([end, left, right], fill="#555555")
    if label:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2 - 22
        draw.text((mx, my), label, font=pil_font(18), fill="#444444", anchor="mm")


def make_architecture(path):
    img = Image.new("RGB", (1800, 1030), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 50), "Kiến trúc local first của V1", font=pil_font(34, True), fill="#111111", anchor="mm")
    box(d, (70, 150, 390, 330), "YouTube", "Video và transcript")
    box(d, (520, 110, 1090, 390), "Chrome Extension", "Sidebar  OAuth  Quiz  Flashcard\nChấm điểm và hiển thị timestamp", "#EAF1F8")
    box(d, (1340, 150, 1730, 330), "Google", "OAuth và Gemini")
    box(d, (520, 570, 1090, 830), "Local RAG Service", "Chunking  Embedding\nRetrieval và Context Builder", "#F5F7FA")
    box(d, (1240, 600, 1690, 800), "ChromaDB cục bộ", "Chunk  Vector  Timestamp")
    arrow(d, (390, 240), (520, 240), "Nội dung video")
    arrow(d, (1090, 220), (1340, 220), "OAuth và AI")
    arrow(d, (1340, 290), (1090, 290), "Kết quả AI")
    arrow(d, (805, 390), (805, 570), "localhost")
    arrow(d, (1090, 700), (1240, 700), "Lưu và truy xuất")
    d.text((900, 950), "OAuth token chỉ được extension gửi tới Google và không đi qua Local RAG Service", font=pil_font(21), fill="#555555", anchor="mm")
    img.save(path)


def make_learning_flow(path):
    img = Image.new("RGB", (1900, 620), "white")
    d = ImageDraw.Draw(img)
    d.text((950, 45), "Luồng học và đánh giá", font=pil_font(34, True), fill="#111111", anchor="mm")
    titles = [
        ("Đọc video", "Transcript\nvà timestamp"),
        ("RAG", "Chunk  Embed\nRetrieve"),
        ("Tạo bài", "Quiz và\nflashcard"),
        ("Người học", "Trả lời\nvà tự đánh giá"),
        ("Chấm điểm", "Theo câu\nvà chủ đề"),
        ("Xem lại", "Điểm mạnh\nđiểm yếu  timestamp"),
    ]
    x = 40
    boxes = []
    for title, subtitle in titles:
        rect = (x, 170, x + 260, 390)
        box(d, rect, title, subtitle, "#F3F6FA" if len(boxes) % 2 == 0 else "#EAF1F8")
        boxes.append(rect)
        x += 315
    for a, b in zip(boxes, boxes[1:]):
        arrow(d, (a[2], 280), (b[0], 280))
    d.text((950, 500), "Câu sai hoặc chủ đề yếu được truy vấn lại trong ChromaDB để tìm đoạn video cần xem", font=pil_font(22), fill="#444444", anchor="mm")
    img.save(path)


make_architecture(ASSETS / "architecture.png")
make_learning_flow(ASSETS / "learning_flow.png")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for name, size, before, after in (("Title", 24, 0, 16), ("Heading 1", 17, 14, 7), ("Heading 2", 13, 10, 4)):
    st = styles[name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Built-in Word Title styles can carry a bottom border in some templates.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("YALA SDS 001   Phiên bản 1.1"), 8.5, color=GRAY)
footer = section.footer.paragraphs[0]
set_font(footer.add_run("YouTube AI Learning Assistant   "), 8.5, color=GRAY)
add_page_number(footer)


def add_para(text="", bold_lead=None, align=None, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, 11, True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, 11)
    else:
        set_font(p.add_run(text), 11)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text), 11)
    return p


def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text), 11)
    return p


def add_table(headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        total_twips = sum(Cm(value).twips for value in widths)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(total_twips))
        tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_ind.set(qn("w:w"), "120")
        layout = tbl_pr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        for i, value in enumerate(widths):
            table.columns[i].width = Cm(value)
            table._tbl.tblGrid.gridCol_lst[i].set(qn("w:w"), str(Cm(value).twips))
    set_table_borders(table)
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(str(text)), font_size, True, "FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_split(table.rows[-1])
        for i, text in enumerate(row):
            cell = cells[i]
            if ridx % 2 == 1:
                shade(cell, ALT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(text)), font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(path, width_cm, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Cm(width_cm))
    set_alt(shape, caption, alt)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(7)
    set_font(cp.add_run(caption), 9.5, color=GRAY, italic=True)


# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(115)
r = p.add_run("SOFTWARE DESIGN SPECIFICATION")
set_font(r, 15, True, NAVY)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_pr = p._p.get_or_add_pPr()
p_border = p_pr.find(qn("w:pBdr"))
if p_border is not None:
    p_pr.remove(p_border)
set_font(p.add_run("YouTube AI Learning Assistant"), 25, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Thiết kế phiên bản V1 cho đồ án môn học"), 13, color=GRAY)
doc.add_paragraph()
add_table(
    ["Thuộc tính", "Giá trị"],
    [
        ("Mã tài liệu", "YALA-SDS-001"),
        ("Phiên bản", "1.1"),
        ("Trạng thái", "Baseline đề nghị phê duyệt"),
        ("Ngày cập nhật", "05/09/2026"),
        ("Phạm vi", "Chrome Extension và Local RAG Service"),
    ],
    [4.4, 11.6],
    10,
)
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Tài liệu này mô tả cấu trúc và hành vi của V1. Các yêu cầu sản phẩm chi tiết được quản lý ngoài SDS."), 10, color=GRAY, italic=True)
doc.add_page_break()


doc.add_heading("Kiểm soát tài liệu", level=1)
add_table(
    ["Phiên bản", "Ngày", "Nội dung", "Trạng thái"],
    [
        ("1.0", "05/09/2026", "Thiết kế kiến trúc local first và luồng học bằng quiz flashcard", "Dự thảo"),
        ("1.1", "05/09/2026", "Bổ sung đặc tả Local REST API, Gemini API, Chrome API và phân công endpoint", "Đề nghị phê duyệt"),
    ],
    [2.5, 3.0, 8.2, 2.5],
)
doc.add_heading("Phê duyệt", level=2)
add_table(
    ["Vai trò", "Họ và tên", "Chữ ký", "Ngày"],
    [
        ("Người lập", "", "", ""),
        ("Trưởng nhóm", "", "", ""),
        ("Giảng viên hoặc người phê duyệt", "", "", ""),
    ],
    [4.2, 4.4, 4.2, 3.0],
)
doc.add_heading("Cách sử dụng tài liệu", level=2)
add_para("SDS xác định cấu trúc phần mềm, trách nhiệm thành phần, luồng dữ liệu, mô hình lưu trữ cục bộ và các giao diện API của V1. Phần API là hợp đồng kỹ thuật giữa Extension, Local RAG Service và Gemini; thay đổi endpoint hoặc schema phải được các bên liên quan review trước khi triển khai.")
doc.add_heading("Mục lục", level=1)
toc_p = doc.add_paragraph()
add_toc(toc_p)
doc.add_page_break()


doc.add_heading("1 Mục đích và phạm vi", level=1)
doc.add_heading("1 1 Mục đích thiết kế", level=2)
add_para("YouTube AI Learning Assistant hỗ trợ một phiên học chủ động trên video YouTube. Extension đọc transcript, Local RAG Service lập chỉ mục nội dung, Gemini tạo quiz và flashcard, sau đó hệ thống chấm điểm và đưa người học về đúng đoạn cần xem lại.")
add_para("Thiết kế V1 ưu tiên một luồng có thể cài đặt và trình diễn trên máy cá nhân. Hệ thống không cần hạ tầng cloud hoặc database nghiệp vụ.")
doc.add_heading("1 2 Phạm vi V1", level=2)
for item in [
    "Chrome Extension chạy trong sidebar trên trang YouTube.",
    "Google OAuth để extension sử dụng Gemini mà không nhúng API key provider.",
    "Thu thập transcript cùng timestamp của video hiện tại.",
    "Chunking, embedding, retrieval và ChromaDB chạy trên máy người dùng.",
    "Tạo quiz, flashcard, chấm điểm và phân tích kết quả học tập.",
    "Đề xuất timestamp cần xem lại trong chính video đang học.",
]:
    add_bullet(item)
doc.add_heading("1 3 Ngoài phạm vi V1", level=2)
for item in [
    "MySQL, tài khoản ứng dụng và server lưu trữ tập trung.",
    "Lịch sử học tập lâu dài hoặc đồng bộ nhiều thiết bị.",
    "Tìm kiếm hoặc tổng hợp kiến thức từ nhiều video.",
    "Gợi ý video mới hoặc xây dựng lộ trình học.",
    "Hỏi đáp tự do với video như một chức năng bắt buộc.",
]:
    add_bullet(item)
doc.add_heading("1 4 Cơ sở tổ chức SDS", level=2)
add_para("Bố cục tham chiếu cách tổ chức Software Design Description trong IEEE 1016-2009 nhưng được điều chỉnh cho quy mô đồ án môn học. Phần yêu cầu chỉ xuất hiện ở mức đầu vào thiết kế; trọng tâm của tài liệu là cấu trúc phần mềm và hành vi khi chạy.")


doc.add_heading("2 Đầu vào và ràng buộc thiết kế", level=1)
add_table(
    ["Nhóm", "Ràng buộc thiết kế"],
    [
        ("Nền tảng", "Chrome Extension Manifest V3 hoạt động trên trang xem video YouTube"),
        ("Xác thực", "Google OAuth; không đọc cookie Google AI Studio và không nhúng API key"),
        ("AI", "Gemini tạo nội dung học tập từ context do RAG cung cấp"),
        ("RAG", "Chunking, embedding và retrieval theo từng video"),
        ("Lưu trữ", "ChromaDB cục bộ chỉ giữ chunk, vector và timestamp metadata"),
        ("Riêng tư", "Không lưu OAuth token, câu trả lời, điểm hoặc hồ sơ học tập lâu dài"),
        ("Triển khai", "Extension và Local RAG Service được cài trên máy người dùng"),
    ],
    [4.0, 12.2],
)
doc.add_heading("2 1 Quyết định phạm vi dữ liệu", level=2)
add_para("Timestamp là metadata của transcript chunk. Embedding hỗ trợ tìm đoạn có liên quan về ngữ nghĩa; ChromaDB nối vector với nội dung và timestamp. Quiz, câu trả lời và kết quả đánh giá chỉ tồn tại trong phiên extension.")


doc.add_heading("3 Kiến trúc tổng thể", level=1)
add_figure(
    ASSETS / "architecture.png",
    16.2,
    "Hình 1 Kiến trúc local first của V1",
    "Sơ đồ kết nối YouTube, Chrome Extension, Google OAuth và Gemini, Local RAG Service và ChromaDB cục bộ.",
)
add_para("Chrome Extension là điểm tương tác của người học. Local RAG Service xử lý transcript và vector trên máy. ChromaDB lưu cache nội dung video cục bộ. Gemini chỉ nhận context cần thiết từ extension và không giao tiếp trực tiếp với ChromaDB.")
doc.add_heading("3 1 Ranh giới trách nhiệm", level=2)
add_table(
    ["Thành phần", "Chịu trách nhiệm", "Không chịu trách nhiệm"],
    [
        ("Chrome Extension", "UI, YouTube, OAuth, phiên quiz, gọi Local API và điều hướng timestamp", "Không lưu vector, chạy embedding hoặc quyết định logic RAG"),
        ("Local RAG Service", "Transcript, chunking, embedding, retrieval, chấm quiz và assessment trong phiên", "Không nhận OAuth token hoặc lưu lịch sử học tập"),
        ("ChromaDB", "Lưu document, embedding và metadata của video", "Không lưu tài khoản, token hoặc lịch sử học"),
        ("Gemini", "Tạo quiz, flashcard và nhận xét từ context", "Không quyết định điểm trắc nghiệm và không truy cập trực tiếp video"),
    ],
    [3.6, 7.3, 5.3],
    8.8,
)


doc.add_heading("4 Thiết kế các thành phần", level=1)
doc.add_heading("4 1 Chrome Extension", level=2)
add_para("Extension gồm content script để quan sát trang YouTube, service worker để điều phối quyền và giao tiếp, cùng sidebar React để hiển thị luồng học. YouTube Adapter cô lập các thao tác phụ thuộc trang như lấy Video ID, đọc timestamp và chuyển player đến vị trí được chọn.")
add_table(
    ["Module", "Trách nhiệm chính"],
    [
        ("YouTube Adapter", "Nhận diện video, transcript, timestamp và thay đổi điều hướng"),
        ("Authentication", "Kiểm tra quyền Google, bắt đầu OAuth sau thao tác người dùng và xử lý lỗi quyền"),
        ("Sidebar", "Hiển thị trạng thái video, quiz, flashcard và báo cáo học tập"),
        ("Quiz Session", "Giữ câu trả lời trong phiên, khóa bài và gửi yêu cầu chấm qua Local API"),
        ("Learning Assessment UI", "Hiển thị điểm, chủ đề mạnh yếu và timestamp do Local API trả về"),
        ("Session Manager", "Quản lý dữ liệu tạm thời và hủy phiên khi đổi video"),
    ],
    [4.2, 12.0],
)
doc.add_heading("4 2 Local RAG Service", level=2)
add_para("Service Python chạy trên localhost và không quản lý người dùng. Nó nhận transcript đã gắn Video ID, tạo hoặc mở index cục bộ, rồi trả về các chunk liên quan cùng timestamp. Service phải có trạng thái sẵn sàng và thao tác xóa cache để extension hướng dẫn người dùng.")
add_table(
    ["Module", "Trách nhiệm chính"],
    [
        ("Transcript Processor", "Chuẩn hóa văn bản nhưng giữ thứ tự và timestamp"),
        ("Chunker", "Tạo các đoạn có kích thước phù hợp và phần chồng lấn"),
        ("Embedding Service", "Tạo vector theo batch bằng model đã chọn"),
        ("Vector Store Adapter", "Tạo collection, thêm dữ liệu, truy vấn và xóa cache ChromaDB"),
        ("Retriever", "Lọc đúng Video ID và xếp hạng chunk liên quan"),
        ("Context Builder", "Sắp xếp chunk, bỏ trùng và giới hạn context gửi sang Gemini"),
    ],
    [4.2, 12.0],
)
doc.add_heading("4 3 ChromaDB cục bộ", level=2)
add_para("V1 dùng ChromaDB PersistentClient để tái sử dụng index sau khi trình duyệt hoặc service khởi động lại. Đây là cache RAG trên thiết bị, không phải database nghiệp vụ. Người dùng có thể xóa toàn bộ cache mà không ảnh hưởng tài khoản Google.")
doc.add_heading("4 4 Gemini", level=2)
add_para("Gemini nhận context transcript đã được giới hạn cùng yêu cầu tạo nội dung. Kết quả phải theo cấu trúc thống nhất để extension kiểm tra trước khi hiển thị. Quiz và flashcard phải gắn được với chủ đề và timestamp nguồn.")


doc.add_heading("5 Thiết kế hành vi khi chạy", level=1)
add_figure(
    ASSETS / "learning_flow.png",
    16.3,
    "Hình 2 Luồng học và đánh giá",
    "Sơ đồ từ đọc transcript, RAG, tạo quiz flashcard, người học trả lời, chấm điểm đến đề xuất timestamp cần xem lại.",
)
doc.add_heading("5 1 Khởi động phiên học", level=2)
for step in [
    "Sidebar đọc Video ID hiện tại và tạo session tương ứng.",
    "Extension kiểm tra quyền Google không tương tác; nếu thiếu quyền thì chờ người dùng bấm đăng nhập.",
    "Extension kiểm tra Local RAG Service và hiển thị hướng dẫn nếu service chưa chạy.",
    "Khi video, xác thực và service sẵn sàng, người dùng có thể lập chỉ mục hoặc sử dụng cache hiện có.",
]:
    add_number(step)
doc.add_heading("5 2 Lập chỉ mục video", level=2)
for step in [
    "Transcript được chuẩn hóa nhưng giữ mốc thời gian gốc.",
    "Chunker tạo các đoạn và gắn start seconds, end seconds cùng thứ tự.",
    "Embedding Service tạo vector cho từng chunk.",
    "Vector Store Adapter ghi document, vector và metadata vào collection.",
    "Service đánh dấu index sẵn sàng cho retrieval.",
]:
    add_number(step)
doc.add_heading("5 3 Tạo nội dung học tập", level=2)
for step in [
    "Retriever chọn các chunk phù hợp với phạm vi bài học.",
    "Context Builder sắp xếp và giới hạn context.",
    "Extension gọi Gemini bằng OAuth và yêu cầu quiz hoặc flashcard.",
    "Extension kiểm tra cấu trúc, đáp án và timestamp trước khi hiển thị.",
]:
    add_number(step)
doc.add_heading("5 4 Chấm điểm và xem lại", level=2)
for step in [
    "Extension gửi quiz và câu trả lời trong phiên tới API assessments quiz.",
    "Assessment Service chấm câu trắc nghiệm và nhóm kết quả theo chủ đề.",
    "Các câu sai hoặc chủ đề yếu được dùng để truy vấn lại ChromaDB.",
    "Gemini tạo nhận xét từ điểm số và chunk được truy xuất.",
    "Sidebar hiển thị điểm mạnh, điểm cần cải thiện và nút nhảy tới timestamp.",
]:
    add_number(step)
doc.add_heading("5 5 Chuyển video", level=2)
add_para("Khi YouTube chuyển sang Video ID khác, extension hủy tác vụ đang chạy, xóa dữ liệu quiz trong phiên và tạo session mới. Cache ChromaDB của video cũ được giữ theo chính sách cục bộ nhưng không được trộn vào truy vấn của video mới.")


doc.add_heading("6 Thiết kế dữ liệu", level=1)
doc.add_heading("6 1 Bản ghi ChromaDB", level=2)
add_table(
    ["Trường", "Ý nghĩa", "Ví dụ"],
    [
        ("id", "Định danh duy nhất của chunk", "videoId chunkIndex"),
        ("document", "Nội dung transcript của chunk", "Đoạn lời nói đã làm sạch"),
        ("embedding", "Vector biểu diễn nội dung", "Mảng số thực"),
        ("video_id", "Video sở hữu chunk", "YouTube Video ID"),
        ("chunk_index", "Thứ tự trong transcript", "0 1 2"),
        ("start_seconds", "Vị trí bắt đầu", "315"),
        ("end_seconds", "Vị trí kết thúc", "352"),
        ("language", "Ngôn ngữ transcript", "vi hoặc en"),
        ("pipeline_version", "Phiên bản chunking và embedding", "1"),
    ],
    [3.3, 7.5, 5.4],
    8.8,
)
doc.add_heading("6 2 Dữ liệu phiên extension", level=2)
add_para("Session Manager chỉ giữ Video ID, trạng thái xử lý, quiz hoặc flashcard hiện tại, câu trả lời và kết quả đánh giá. Dữ liệu này không được ghi vào ChromaDB. Khi phiên kết thúc, dữ liệu có thể bị loại bỏ mà không cần quy trình migration.")
doc.add_heading("6 3 Chính sách cache", level=2)
for item in [
    "Index được tái sử dụng khi Video ID và pipeline version trùng khớp.",
    "Người dùng có thể xóa cache từ giao diện hoặc script cục bộ.",
    "Khi thay đổi embedding model hoặc chunking, pipeline version phải tăng.",
    "V1 không tự động suy luận danh tính người dùng từ dữ liệu cache.",
]:
    add_bullet(item)


doc.add_heading("7 Thiết kế giao diện và API", level=1)
add_para("V1 có ba nhóm giao diện lập trình: Chrome Extension API trong trình duyệt, Local REST API giữa Extension và Local RAG Service, và Gemini API bên ngoài. Local REST API được version hóa bằng tiền tố api v1, trao đổi JSON qua loopback và không nhận OAuth token.")

doc.add_heading("7 1 Tổng quan giao tiếp", level=2)
add_table(
    ["Nguồn", "Đích", "Giao diện", "Mục đích"],
    [
        ("YouTube Adapter", "Extension", "Chrome runtime và message", "Video ID, transcript, timestamp và điều khiển player"),
        ("Extension", "Local RAG Service", "HTTP JSON trên 127.0.0.1", "Index, retrieval, assessment và xóa cache"),
        ("Local RAG Service", "ChromaDB", "Python client cục bộ", "Lưu và truy vấn chunk, embedding, metadata"),
        ("Extension", "Google", "OAuth 2.0 và Gemini REST", "Xác thực và tạo quiz, flashcard, nhận xét"),
    ],
    [3.0, 3.6, 5.3, 4.3],
    8.3,
)

doc.add_heading("7 2 Quy ước Local REST API", level=2)
for item in [
    "Base URL mặc định là http://127.0.0.1 với port được cấu hình; service không bind ra mạng ngoài theo mặc định.",
    "Tất cả endpoint V1 dùng tiền tố /api/v1 và Content Type application/json, trừ response không có nội dung nếu được quy định.",
    "Tên trường JSON dùng camelCase tại biên API; module Python được phép mapping sang snake case ở lớp transport.",
    "Mỗi request liên quan video phải có videoId khớp giữa path và payload; sai khác bị từ chối.",
    "OAuth token chỉ được Extension gửi tới Google. Local REST API không có trường accessToken và không chấp nhận Authorization Bearer của Google.",
    "Transport giới hạn kích thước transcript, kiểm tra Origin extension theo cấu hình và không ghi toàn bộ transcript vào log.",
]:
    add_bullet(item)

add_table(
    ["Method", "Endpoint", "Chức năng", "Chủ sở hữu"],
    [
        ("GET", "/api/v1/health", "Kiểm tra service và dependency", "TV2"),
        ("POST", "/api/v1/videos/{videoId}/index", "Lập hoặc tái sử dụng index", "HTTP TV2; logic TV3"),
        ("GET", "/api/v1/videos/{videoId}/index-status", "Đọc trạng thái index", "HTTP TV2; logic TV3"),
        ("POST", "/api/v1/videos/{videoId}/retrieve", "Lấy context cho quiz flashcard review", "HTTP TV2; logic TV3"),
        ("POST", "/api/v1/videos/{videoId}/assessments/quiz", "Chấm quiz và tìm đoạn cần xem", "HTTP TV2; logic TV3"),
        ("DELETE", "/api/v1/videos/{videoId}/cache", "Xóa cache của video", "HTTP TV2; logic TV3"),
    ],
    [2.0, 6.1, 4.8, 3.3],
    7.9,
)

doc.add_heading("7 3 GET api v1 health", level=2)
add_table(
    ["Thuộc tính", "Đặc tả"],
    [
        ("Mục đích", "Cho Extension biết Local RAG Service có sẵn sàng nhận yêu cầu hay không."),
        ("Request", "Không có body."),
        ("Response 200", "status, serviceVersion, pipelineVersion, vectorStoreReady và embeddingModelReady."),
        ("Response 503", "status bằng not_ready và error mô tả dependency chưa sẵn sàng."),
        ("Hiện thực", "TV2 code route, lifecycle và health aggregation; TV3 cung cấp trạng thái embedding và ChromaDB."),
        ("Kiểm thử", "TV4 kiểm tra service sẵn sàng, thiếu model, ChromaDB lỗi và service offline."),
    ],
    [3.8, 12.4],
    8.7,
)

doc.add_heading("7 4 POST api v1 videos videoId index", level=2)
add_table(
    ["Thuộc tính", "Đặc tả"],
    [
        ("Mục đích", "Nhận video và transcript có timestamp, sau đó chuẩn hóa, chunk, embed và lưu ChromaDB; dùng lại cache nếu cùng pipelineVersion."),
        ("Path", "videoId bắt buộc và phải khớp payload.video.videoId."),
        ("Request", "video gồm videoId, title, durationSec, language; transcriptSegments gồm text, startSec, endSec, position."),
        ("Response 200", "indexStatus bằng ready và cached bằng true khi index phù hợp đã tồn tại."),
        ("Response 202", "videoId, indexStatus bằng indexing, cached bằng false và pipelineVersion khi bắt đầu xử lý."),
        ("Lỗi", "400 INVALID_REQUEST; 409 VIDEO_ID_MISMATCH; 413 PAYLOAD_TOO_LARGE; 422 TRANSCRIPT_INVALID; 500 INDEX_FAILED; 503 SERVICE_NOT_READY."),
        ("Hiện thực", "TV2 code route, validation, giới hạn payload và response; TV3 code transcript processor, chunking, embedding và vector store."),
        ("Kiểm thử", "TV4 kiểm tra transcript hợp lệ, rỗng, timestamp sai, request quá lớn, cached index và hai video độc lập."),
    ],
    [3.8, 12.4],
    8.6,
)

doc.add_heading("7 5 GET api v1 videos videoId index status", level=2)
add_table(
    ["Thuộc tính", "Đặc tả"],
    [
        ("Mục đích", "Cho Extension theo dõi hoặc kiểm tra trạng thái index của video hiện tại."),
        ("Request", "Không có body; videoId nằm trên path."),
        ("Response 200", "videoId, indexStatus thuộc not_indexed indexing ready failed, chunkCount, pipelineVersion và error nếu có."),
        ("Lỗi", "400 INVALID_VIDEO_ID; 500 VECTOR_STORE_ERROR; 503 SERVICE_NOT_READY."),
        ("Hiện thực", "TV2 code route và response; TV3 cung cấp index repository và trạng thái collection."),
        ("Kiểm thử", "TV4 kiểm tra video chưa index, đang index, hoàn thành, thất bại và restart service."),
    ],
    [3.8, 12.4],
    8.7,
)

doc.add_heading("7 6 POST api v1 videos videoId retrieve", level=2)
add_table(
    ["Thuộc tính", "Đặc tả"],
    [
        ("Mục đích", "Truy xuất chunk cùng timestamp làm context cho quiz, flashcard hoặc review."),
        ("Request", "query, purpose thuộc quiz flashcard review, và maxResults tùy chọn; service áp dụng giới hạn cấu hình."),
        ("Response 200", "videoId, purpose, chunks; mỗi chunk có chunkId, text, startSec, endSec, score và position."),
        ("Không đủ căn cứ", "Trả chunks rỗng và reason bằng NO_RELEVANT_CONTEXT thay vì tạo nội dung không bám transcript."),
        ("Lỗi", "400 INVALID_REQUEST; 404 INDEX_NOT_FOUND; 422 QUERY_INVALID; 500 RETRIEVAL_FAILED; 503 SERVICE_NOT_READY."),
        ("Hiện thực", "TV2 code route và mapping; TV3 code query embedding, filter videoId, ranking, threshold, deduplication và Context Builder."),
        ("Kiểm thử", "TV4 dùng expected chunks và timestamps của TV3; bắt buộc kiểm tra không trộn dữ liệu giữa video."),
    ],
    [3.8, 12.4],
    8.6,
)

doc.add_heading("7 7 POST api v1 videos videoId assessments quiz", level=2)
add_table(
    ["Thuộc tính", "Đặc tả"],
    [
        ("Mục đích", "Chấm bài trong phiên, nhóm chủ đề mạnh yếu và tìm timestamp cần xem lại."),
        ("Request", "quizId tùy chọn, questions đã validate gồm đáp án và topic, cùng userAnswers; không chứa OAuth token hoặc thông tin tài khoản."),
        ("Response 200", "score, correctCount, totalCount, questionResults, strongTopics, weakTopics và reviewTimestamps có topic, reason, startSec, endSec, chunkId."),
        ("Lỗi", "400 INVALID_REQUEST; 404 INDEX_NOT_FOUND; 422 QUIZ_INVALID; 500 ASSESSMENT_FAILED; 503 SERVICE_NOT_READY."),
        ("Lưu trữ", "Request và response chỉ tồn tại trong phiên xử lý; không ghi câu trả lời, điểm hoặc assessment vào ChromaDB."),
        ("Hiện thực", "TV2 code route và validation; TV3 code scoring, topic analyzer, review retrieval và kết quả assessment."),
        ("Kiểm thử", "TV4 kiểm tra điểm, chủ đề mạnh yếu, câu bỏ trống, timestamp review và không lưu dữ liệu người học."),
    ],
    [3.8, 12.4],
    8.5,
)

doc.add_heading("7 8 DELETE api v1 videos videoId cache", level=2)
add_table(
    ["Thuộc tính", "Đặc tả"],
    [
        ("Mục đích", "Xóa chunk, embedding và metadata của đúng một video khỏi cache cục bộ."),
        ("Request", "Không có body; videoId nằm trên path."),
        ("Response 200", "videoId, deleted và deletedChunkCount; gọi lại an toàn với deleted bằng false."),
        ("Lỗi", "400 INVALID_VIDEO_ID; 500 CACHE_DELETE_FAILED; 503 SERVICE_NOT_READY."),
        ("Hiện thực", "TV2 code route và xác nhận đường dẫn vận hành; TV3 code thao tác xóa collection record theo videoId."),
        ("Kiểm thử", "TV4 xác minh chỉ dữ liệu video được chọn bị xóa và video khác vẫn truy vấn được."),
    ],
    [3.8, 12.4],
    8.7,
)

doc.add_heading("7 9 Cấu trúc lỗi chung", level=2)
add_para("Response lỗi dùng đối tượng error gồm code, message, retryable và details an toàn. Message dành cho người dùng không chứa stack trace, đường dẫn nhạy cảm, transcript đầy đủ hoặc credential.")
add_table(
    ["HTTP", "Mã lỗi tiêu biểu", "Ý nghĩa"],
    [
        ("400", "INVALID_REQUEST hoặc INVALID_VIDEO_ID", "Request không đúng contract"),
        ("404", "INDEX_NOT_FOUND", "Video chưa có index phù hợp"),
        ("409", "VIDEO_ID_MISMATCH", "Video ID trên path khác payload"),
        ("413", "PAYLOAD_TOO_LARGE", "Transcript vượt giới hạn"),
        ("422", "TRANSCRIPT_INVALID QUIZ_INVALID QUERY_INVALID", "Dữ liệu đúng JSON nhưng không hợp lệ nghiệp vụ"),
        ("500", "INDEX_FAILED RETRIEVAL_FAILED ASSESSMENT_FAILED", "Lỗi xử lý nội bộ đã được che thông tin nhạy cảm"),
        ("503", "SERVICE_NOT_READY", "Model hoặc vector store chưa sẵn sàng"),
    ],
    [2.0, 6.0, 8.2],
    8.3,
)

doc.add_heading("7 10 Gemini API", level=2)
add_para("Extension gọi trực tiếp phương thức models generateContent tại https://generativelanguage.googleapis.com/v1beta/{model=models/*}:generateContent. Request dùng OAuth access token trong Authorization Bearer, project/quota header theo cấu hình Google, Content Type application/json và context đã được Local RAG Service giới hạn.")
add_table(
    ["Thuộc tính", "Thiết kế"],
    [
        ("Chức năng", "Tạo quiz, flashcard và nhận xét học tập có cấu trúc."),
        ("Request", "systemInstruction, contents chứa context RAG và task, generationConfig, response MIME type JSON và response schema phù hợp."),
        ("Response", "Candidates được chuyển đến ai content validation; response thô không đi thẳng vào UI."),
        ("Lỗi", "401 yêu cầu làm mới OAuth; 403 thiếu quyền/project; 429 hết quota; 5xx lỗi Google có thể retry giới hạn."),
        ("Ownership", "TV1 code OAuth và Gemini transport; TV3 code prompt, schema, validator, mapper; TV4 kiểm thử auth, quota và output sai."),
    ],
    [3.8, 12.4],
    8.6,
)
add_para("Đăng nhập trang Google AI Studio không được dùng như cơ chế xác thực API. V1 phải có Google Cloud project, bật Generative Language API, cấu hình OAuth consent và OAuth client cho ứng dụng. Extension không đọc cookie AI Studio và không nhúng provider API key.")

doc.add_heading("7 11 Chrome Extension API", level=2)
add_table(
    ["Chrome API", "Mục đích", "Owner"],
    [
        ("chrome identity", "Kiểm tra token không tương tác; chỉ mở OAuth tương tác sau thao tác người dùng", "TV1; TV4 test"),
        ("chrome sidePanel", "Hiển thị giao diện học bên cạnh trang YouTube", "TV1"),
        ("chrome runtime", "Message giữa service worker, content script và side panel", "TV1"),
        ("chrome tabs", "Nhận diện tab và gửi message tới content script", "TV1"),
        ("content script", "Đọc video, transcript, timestamp và điều khiển player qua YouTube Adapter", "TV1; TV4 regression"),
    ],
    [4.0, 8.7, 3.5],
    8.4,
)


doc.add_heading("8 Thiết kế trạng thái và lỗi", level=1)
doc.add_heading("8 1 Trạng thái chính", level=2)
add_para("BOOTSTRAPPING -> CHECKING AUTH -> CHECKING LOCAL SERVICE -> READING VIDEO -> INDEXING -> READY -> GENERATING -> LEARNING -> GRADING -> RESULT")
add_para("Các trạng thái lỗi có thể quay lại bước phù hợp mà không tải lại toàn bộ trang: AUTH REQUIRED, PERMISSION DENIED, LOCAL SERVICE OFFLINE, TRANSCRIPT UNAVAILABLE, INDEXING FAILED, AI UNAVAILABLE và INVALID AI RESULT.")
doc.add_heading("8 2 Phản ứng với lỗi", level=2)
add_table(
    ["Tình huống", "Phản ứng thiết kế"],
    [
        ("Chưa có quyền Google", "Hiển thị nút đăng nhập; không tự bật cửa sổ OAuth"),
        ("Local RAG Service chưa chạy", "Hiển thị hướng dẫn khởi động và nút kiểm tra lại"),
        ("Video không có transcript", "Dừng lập chỉ mục và giải thích lý do"),
        ("Embedding hoặc ChromaDB lỗi", "Giữ trạng thái video và cho phép thử lập chỉ mục lại"),
        ("Gemini trả sai cấu trúc", "Không hiển thị nội dung chưa kiểm tra và cho phép tạo lại"),
        ("Người dùng chuyển video", "Hủy tác vụ cũ và ngăn kết quả cũ xuất hiện ở video mới"),
    ],
    [5.1, 11.1],
)


doc.add_heading("9 Tổ chức mã nguồn", level=1)
add_para("Repository tách extension khỏi Local RAG Service nhưng đặt schema và tài liệu chung trong cùng một kho để bốn thành viên phối hợp.")
code = [
    "youtube-ai-learning-assistant/",
    "|-- extension/",
    "|   |-- public/",
    "|   `-- src/",
    "|       |-- background/",
    "|       |-- content/",
    "|       |-- sidebar/",
    "|       |-- features/",
    "|       |-- integrations/",
    "|       |   |-- local-service/",
    "|       |   |-- google-oauth/",
    "|       |   `-- gemini/",
    "|       |-- ai-content/",
    "|       |   |-- prompts/",
    "|       |   |-- validation/",
    "|       |   `-- mappers/",
    "|       |-- session/",
    "|       `-- types/",
    "|-- local-rag-service/",
    "|   |-- app/",
    "|   |   |-- core/",
    "|   |   |-- transport/",
    "|   |   |-- transcript/",
    "|   |   |-- chunking/",
    "|   |   |-- embedding/",
    "|   |   |-- vector_store/",
    "|   |   |-- retrieval/",
    "|   |   `-- assessment/",
    "|   `-- tests/",
    "|-- shared/contracts/",
    "|-- scripts/",
    "|-- tests/e2e/",
    "|-- docs/architecture/",
    "|-- docs/project-management/",
    "|-- docs/testing/",
    "`-- README.md",
]
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(0.7)
for i, line in enumerate(code):
    r = p.add_run(line + ("\n" if i < len(code) - 1 else ""))
    r.font.name = "Consolas"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(9.3)
    r.font.color.rgb = RGBColor(0, 0, 0)
doc.add_heading("9 1 Quy tắc phụ thuộc", level=2)
for item in [
    "Sidebar không truy cập trực tiếp DOM YouTube; mọi thao tác đi qua YouTube Adapter.",
    "Module AI không gọi trực tiếp ChromaDB; context đến từ Local RAG Service.",
    "Vector Store Adapter là nơi duy nhất phụ thuộc ChromaDB.",
    "Prompt được đặt trong thư mục riêng và không viết rải rác trong component.",
    "Schema dùng chung được cập nhật trước khi thay đổi cả extension và service.",
]:
    add_bullet(item)


doc.add_heading("10 Bảo mật và riêng tư", level=1)
add_table(
    ["Khu vực", "Quy tắc thiết kế"],
    [
        ("OAuth", "Token chỉ do extension quản lý và chỉ gửi tới Google"),
        ("Cookie", "Extension không khai báo quyền để đọc cookie AI Studio"),
        ("Local service", "Chỉ lắng nghe trên loopback và chỉ nhận dữ liệu cần cho RAG"),
        ("ChromaDB", "Không lưu token, email, câu trả lời, điểm hoặc nhận xét cá nhân"),
        ("Log", "Không ghi credential hoặc toàn bộ dữ liệu nhạy cảm"),
        ("Xóa dữ liệu", "Có thao tác xóa cache ChromaDB trên máy người dùng"),
    ],
    [4.0, 12.2],
)
add_para("Content script được xem là nguồn dữ liệu cần kiểm tra. Service worker và Local RAG Service phải xác thực Video ID, kích thước transcript và kiểu dữ liệu trước khi xử lý.")


doc.add_heading("11 Cài đặt và vận hành V1", level=1)
doc.add_heading("11 1 Cài đặt", level=2)
for step in [
    "Cài môi trường cần thiết cho Local RAG Service bằng script của repository.",
    "Khởi động service cục bộ bằng start local service bat.",
    "Build extension và chọn Load unpacked trong Chrome.",
    "Mở video YouTube có transcript rồi mở sidebar.",
]:
    add_number(step)
doc.add_heading("11 2 Dữ liệu trên máy", level=2)
add_para("ChromaDB được đặt trong thư mục dữ liệu cục bộ đã công bố trong README. V1 phải có script clear local data để người dùng xóa index. Việc gỡ extension không được mô tả là tự động xóa thư mục ChromaDB nếu chưa có cơ chế cài đặt quản lý thư mục đó.")
doc.add_heading("11 3 Hướng phát triển đóng gói", level=2)
add_para("Sau V1, Local RAG Service có thể được đóng gói thành executable hoặc Native Messaging Host để giảm thao tác khởi động. Quyết định này không ảnh hưởng mô hình dữ liệu hoặc luồng RAG của V1.")


doc.add_heading("12 Xác minh thiết kế", level=1)
add_table(
    ["Nội dung cần chứng minh", "Bằng chứng dự kiến"],
    [
        ("YouTube integration", "Demo nhận diện video, đọc timestamp và chuyển player"),
        ("Google OAuth", "Test đăng nhập mới, đã cấp quyền, từ chối quyền và hết hạn"),
        ("RAG", "Bộ truy vấn kiểm thử trả đúng chunk thuộc video và timestamp"),
        ("ChromaDB cục bộ", "Kiểm tra index tồn tại trên máy và có thể xóa"),
        ("Quiz và flashcard", "Kết quả AI hợp lệ và có nguồn transcript"),
        ("Chấm điểm", "Unit test tính điểm và nhóm chủ đề"),
        ("Đánh giá học tập", "Câu sai dẫn tới đúng phần transcript cần xem lại"),
        ("Riêng tư", "Không tìm thấy token, email hoặc điểm trong ChromaDB và log"),
    ],
    [6.0, 10.2],
)
doc.add_heading("12 1 Luồng nghiệm thu", level=2)
add_para("Mở video -> mở extension -> xác thực Google -> lập chỉ mục transcript -> tạo quiz -> làm bài -> chấm điểm -> hiển thị phần học tốt và phần cần cải thiện -> bấm timestamp để xem lại đúng đoạn video.")


doc.add_heading("13 Quyết định thiết kế và rủi ro", level=1)
doc.add_heading("13 1 Quyết định chính", level=2)
add_table(
    ["Mã", "Quyết định", "Lý do"],
    [
        ("D01", "Dùng Local RAG Service", "ChromaDB và embedding chạy cục bộ mà không cần server Internet"),
        ("D02", "Dùng ChromaDB PersistentClient", "Tái sử dụng index video sau khi khởi động lại"),
        ("D03", "Không dùng MySQL trong V1", "Không có dữ liệu nghiệp vụ hoặc tài khoản cần quản lý"),
        ("D04", "Token chỉ ở extension", "Giữ ranh giới xác thực Google tách khỏi local service"),
        ("D05", "Assessment chạy trong Local RAG Service", "Thành viên AI RAG sở hữu logic; request và kết quả chỉ tồn tại trong phiên, không lưu lâu dài"),
    ],
    [1.5, 6.3, 8.4],
    8.7,
)
doc.add_heading("13 2 Rủi ro cần kiểm chứng sớm", level=2)
add_table(
    ["Rủi ro", "Cách giảm thiểu"],
    [
        ("Không lấy được transcript ổn định", "Prototype với video tiếng Việt, tiếng Anh, ngắn, dài và thiếu subtitle"),
        ("OAuth hoặc quyền Gemini khác giả định", "Thực hiện proof of concept trước khi xây chức năng AI"),
        ("Model embedding quá nặng", "Benchmark model đa ngôn ngữ nhỏ trên máy cấu hình trung bình"),
        ("Cache tăng nhanh", "Giới hạn dung lượng và cung cấp thao tác xóa cache"),
        ("AI tạo câu hỏi ngoài transcript", "Giới hạn context và kiểm tra timestamp nguồn"),
        ("YouTube thay đổi giao diện", "Cô lập logic trong YouTube Adapter và duy trì regression test"),
    ],
    [6.2, 10.0],
)


doc.add_heading("14 Trách nhiệm triển khai", level=1)
add_table(
    ["Thành viên", "Phạm vi thiết kế", "Đầu ra chính"],
    [
        ("1", "Extension, sidebar, YouTube Adapter và OAuth UI", "Bản extension load được và các màn hình V1"),
        ("2", "Local Service core, REST transport và vận hành cục bộ", "Endpoint chạy được, health, validation và script vận hành"),
        ("3", "Transcript, chunking, embedding, ChromaDB, retrieval, prompt và assessment", "Pipeline AI RAG, logic sau endpoint và bộ đánh giá"),
        ("4", "Integration, bảo mật, test và tài liệu", "Test report, README, slide và demo"),
    ],
    [2.2, 8.0, 6.0],
    8.8,
)


doc.add_heading("15 Tài liệu tham khảo", level=1)
for item in [
    "IEEE 1016-2009 Software Design Descriptions",
    "Chrome Extensions Manifest V3 và Chrome Identity API",
    "Chrome Extension Storage và Native Messaging documentation",
    "Chroma Clients và PersistentClient documentation",
    "Google Gemini OAuth documentation https://ai.google.dev/gemini-api/docs/oauth",
    "Google Gemini models generateContent API https://ai.google.dev/api/generate-content",
]:
    add_bullet(item)


doc.add_page_break()
doc.add_heading("Đề nghị phê duyệt", level=1)
add_para("Nhóm đề nghị sử dụng tài liệu này làm baseline thiết kế cho V1. Những thay đổi làm xuất hiện server Internet, database nghiệp vụ, lịch sử học tập hoặc gợi ý video phải được xem là thay đổi phạm vi và được nhóm phê duyệt trước khi triển khai.")
add_table(
    ["Phương án", "Đánh dấu", "Ý kiến"],
    [
        ("Phê duyệt thiết kế V1", "", ""),
        ("Phê duyệt có điều kiện", "", ""),
        ("Yêu cầu chỉnh sửa", "", ""),
    ],
    [7.0, 2.3, 6.9],
)


settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.core_properties.title = "Software Design Specification YouTube AI Learning Assistant"
doc.core_properties.subject = "Thiết kế V1 local first với Chrome Extension Local RAG Service và ChromaDB cục bộ"
doc.core_properties.author = "Nhóm dự án YouTube AI Learning Assistant"
doc.core_properties.keywords = "SDS, Chrome Extension, RAG, ChromaDB, Gemini, YouTube"
doc.save(OUT)
print(OUT)
